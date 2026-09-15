from io import BytesIO
import PyPDF2
from docx import Document
from bs4 import BeautifulSoup
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_pinecone import PineconeVectorStore
import os
from langchain_core.documents import Document
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_google_genai.embeddings import GoogleGenerativeAIEmbeddings
from dotenv import load_dotenv
from typing import Annotated, List
from pinecone import Pinecone, ServerlessSpec
from langchain_core.messages import HumanMessage, AIMessage, SystemMessage
from textwrap import dedent
load_dotenv()
google_api_key = os.environ.get("api_key")
pinecone_key = os.environ.get("PINECONE_API_KEY")

def create_index_if_not_exists():
    try:
        pc = Pinecone(api_key=pinecone_key)
        if "rag-index" not in pc.list_indexes().names():
            pc.create_index(
                name="rag-index",
                dimension=3072,  # For gemini-embedding-2
                metric="cosine",
                spec=ServerlessSpec(
                    cloud="aws",
                    region="us-east-1",
                )
            )
            print("Index 'rag-index' created successfully")
        else:
            print("Index 'rag-index' already exists")
    except Exception as e:
        raise ValueError(f"Error creating index: {str(e)}")

if not google_api_key or not pinecone_key:
    raise ValueError("One or more required API keys not found in environment variables. Please set 'api_key' and 'PINECONE_API_KEY' in your .env file.")
try:
    
    embedder = GoogleGenerativeAIEmbeddings(model="gemini-embedding-2",api_key=google_api_key)
    llm_model = ChatGoogleGenerativeAI(model="gemini-3.1-flash-lite", api_key=google_api_key, temperature=0.2)
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=100,
        separators=["\n\n", "\n", " ", ""]
    )
    VectorStore = PineconeVectorStore(
        pinecone_api_key=pinecone_key,
        index_name="rag-index",
        embedding=embedder
    )
except Exception as e:
    raise ValueError(f"Error initializing PineconeVectorStore: {str(e)}")

def extract_text_from_pdf(file: bytes, filename: str) -> dict:
    try:
        pdf_reader = PyPDF2.PdfReader(BytesIO(file))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        if not text.strip():
            raise ValueError("No text found in the PDF. It might be a scanned document or image-based PDF.")
        chunks = text_splitter.split_text(text)
        documents = [
            Document(page_content=chunk, metadata={"source": filename, "chunk_index": i})
            for i, chunk in enumerate(chunks)
            ]
        VectorStore.add_documents(documents)
        return {"status":"success", "message": f"PDF processed and stored in vector store. Total chunks: {len(chunks)}"}
    except Exception as e:
        raise ValueError(f"Error extracting text from PDF: {str(e)}")

def extract_text_from_docx(file: bytes, filename: str) -> dict:
    try:
        doc = Document(BytesIO(file))
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        if not text.strip():
            raise ValueError("No text found in the DOCX file.")
        chunks = text_splitter.split_text(text)
        documents = [
        Document(page_content=chunk, metadata={"source": filename, "chunk_index": i})
        for i, chunk in enumerate(chunks)
        ]
        VectorStore.add_documents(documents)
        return {"status":"success", "message": f"DOCX processed and stored in vector store. Total chunks: {len(chunks)}"}
    except Exception as e:
        raise ValueError(f"Error extracting text from DOCX: {str(e)}")

def extract_text_from_file(file: bytes, filename: str, extension: str) -> dict:
    if extension=="pdf":
        return extract_text_from_pdf(file, filename)
    if extension=="docx":
        return extract_text_from_docx(file, filename)
    if extension=="txt":
        try:
            text = file.decode("utf-8")
            if not text.strip():
                raise ValueError("No text found in the TXT file.")
            chunks = text_splitter.split_text(text)
            documents = [
                Document(page_content=chunk, metadata={"source": filename, "chunk_index": i})
                for i, chunk in enumerate(chunks)
            ]
            VectorStore.add_documents(documents)
            return {"status":"success", "message": f"TXT processed and stored in vector store. Total chunks: {len(chunks)}"}
        except Exception as e:
            raise ValueError(f"Error extracting text from TXT: {str(e)}")
    if extension=="html":
        try:
            text = file.decode("utf-8")
            soup = BeautifulSoup(text, "html.parser")
            text = soup.get_text().strip()
            if not text:
                raise ValueError("No text found in the HTML file.")
            chunks = text_splitter.split_text(text)
            documents = [
                Document(page_content=chunk, metadata={"source": filename, "chunk_index": i})
                for i, chunk in enumerate(chunks)
            ]
            VectorStore.add_documents(documents)
            return {"status":"success", "message": f"HTML processed and stored in vector store. Total chunks: {len(chunks)}"}
        except Exception as e:
            raise ValueError(f"Error extracting text from HTML: {str(e)}")
    if extension=="doc":
        raise ValueError("DOC files are not supported. Please convert to DOCX.")
    raise ValueError(f"Unsupported file extension: {extension}")

def rag_chat(query: str, messages: List):
    try:
        relevant_docs = VectorStore.similarity_search_with_score(query, k=3)
        filtered_docs = [
            doc for doc, score in relevant_docs
            if score > 0.6
        ]
        if not filtered_docs:
            yield "I don't have enough context to answer this question. Please try rephrasing your question."
            return
        context = "\n\n".join([doc.page_content for doc in filtered_docs])

        system_message = [SystemMessage(content=dedent('''You are a helpful assistant that provides information based on the provided context.
                     Use the context to answer the user's questions.
                     If the answer is not in the context, respond with "I don't know." Do not make up answers.
                     Be concise and accurate.
                     Stick to the facts provided in the context. Do not provide any information that is not present in the context.
                     Do not perform any other action other than answering the user's questions based on the context.
                     Context: {context}''').format(context=context))]

        for chunk in llm_model.stream(system_message + messages + [HumanMessage(content=query)]):
            if hasattr(chunk, 'content') and isinstance(chunk.content, list):
                for item in chunk.content:
                    if isinstance(item, dict) and 'text' in item and item['text']:
                        yield item['text']
    except ValueError as e:
        raise ValueError(f"Vector store error: {str(e)}")
    except Exception as e:
        raise ValueError(f"Error during RAG chat: {str(e)}")